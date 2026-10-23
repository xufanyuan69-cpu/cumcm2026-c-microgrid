import openpyxl
import numpy as np
from scipy.optimize import linprog
from pathlib import Path

from project_paths import input_directory, analysis_directory, prepare_result

base = input_directory()
ws = openpyxl.load_workbook(base/'附件1.xlsx', data_only=True).active
price_raw = np.array([ws.cell(i,2).value for i in range(2,146)], float)
load_raw = np.array([ws.cell(i,3).value for i in range(2,146)], float)
pv_raw = np.array([ws.cell(i,4).value for i in range(2,146)], float)
T=len(price_raw); dt=1/6; eta=.9; pmax=5000*dt
# 附件按 0:10,...,23:50,次日0:00 排列；SOC 却以 0:00 为边界。
# 因此把末行的 0:00 循环移到首位，按真实 0:00-24:00 顺序优化。
chrono_to_raw=np.r_[T-1,np.arange(T-1)]
price=price_raw[chrono_to_raw]
load=load_raw[chrono_to_raw]
pv=pv_raw[chrono_to_raw]
def ix(block,t): return block*T+t
n=5*T; obj=np.zeros(n); obj[:T]=price
Aeq=[]; beq=[]
for t in range(T):
    row=np.zeros(n); row[ix(0,t)]=1; row[ix(1,t)]=-1; row[ix(2,t)]=1; row[ix(3,t)]=-1
    Aeq.append(row); beq.append((load[t]-pv[t])*dt)
    row=np.zeros(n); row[ix(4,t)]=1; row[ix(1,t)]=-eta; row[ix(2,t)]=1/eta
    rhs=6000
    if t>0: row[ix(4,t-1)]-=1; rhs=0
    Aeq.append(row); beq.append(rhs)
row=np.zeros(n); row[ix(4,T-1)]=1; Aeq.append(row); beq.append(6000)
bounds=[(0,None)]*n
for t in range(T):
    bounds[ix(1,t)]=(0,pmax)
    bounds[ix(2,t)]=(0,pmax)
    bounds[ix(3,t)]=(0,pv[t]*dt)
    bounds[ix(4,t)]=(1200,10800)
# 第一阶段：严格最小化购电费用。
Aeq=np.array(Aeq); beq=np.array(beq)
stage1=linprog(obj,A_eq=Aeq,b_eq=beq,bounds=bounds,method='highs')
if not stage1.success: raise RuntimeError(stage1.message)

# 第二阶段：在最优购电费用几乎不变的条件下，最小化电池吞吐量。
# 这比直接把任意 epsilon 加进原目标更稳健，也能消除无意义的循环充放电。
throughput=np.zeros(n); throughput[T:3*T]=1.0
res=linprog(
    throughput,
    A_eq=np.vstack([Aeq,obj]),
    b_eq=np.append(beq,stage1.fun),
    bounds=bounds,
    method='highs',
)
if not res.success: raise RuntimeError(res.message)
x=res.x; P=x[:T]; C=x[T:2*T]; D=x[2*T:3*T]; S=x[4*T:5*T]
for arr in (P,C,D,S,x[3*T:4*T]): arr[np.abs(arr)<1e-8]=0.0
actual_cost=float(price@P)
simultaneous=int(np.sum((C>1e-7)&(D>1e-7)))
print('objective',actual_cost,'Ptotal',P.sum(),'throughput',C.sum()+D.sum())
print('max residual',max(abs(Aeq@x-beq)),'simultaneous intervals',simultaneous)
print('SOC range',S.min(),S.max(),'terminal',S[-1])
out=prepare_result('result1.xlsx'); wb=openpyxl.load_workbook(out)
wsp=wb['计划购电量']; wsc=wb['充放电量']
P_raw=np.empty(T); P_raw[chrono_to_raw]=P
for t in range(T): wsp.cell(t+2,2).value=float(P_raw[t])
for g in range(6):
    a,b=24*g,24*(g+1)
    wsc.cell(g+2,2).value=float(C[a:b].sum()); wsc.cell(g+2,3).value=float(D[a:b].sum())
wsc.cell(2,5).value=6000.0; wsc.cell(3,5).value=float(S[-1])
wb.save(out); print('saved',out)
analysis_dir=analysis_directory()
np.savez(
    analysis_dir/'q1_solution.npz',
    price=price,load=load,pv=pv,P=P,C=C,D=D,R=x[3*T:4*T],S=S,
)
for label in ['10:00-10:10','12:00-12:10','14:00-14:10','16:00-16:10','18:00-18:10','20:00-20:10']:
    row=next(r for r in range(2,146) if wsp.cell(r,1).value==label)
    print('table1',label,wsp.cell(row,2).value)
print('table1_total',P.sum(),'cost',actual_cost)
for g in range(6): print('table2',g,C[24*g:24*(g+1)].sum(),D[24*g:24*(g+1)].sum())
