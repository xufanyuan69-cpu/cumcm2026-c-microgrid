from pathlib import Path
from copy import copy
import numpy as np
import openpyxl
from scipy.optimize import linprog
import pandas as pd
from dispatch_execution import execute_fixed_plan

import os
from project_paths import input_directory, analysis_directory, prepare_result

BASE=input_directory()
OUT=analysis_directory()
T=144; DT=1/6; ETA=.9; PMAX=5000*DT; SMIN=1200.; SMAX=10800.
ALPHA=.90; WINDOW=30; RISK_AWARE=os.environ.get('CUMCM_RISK_AWARE', '1') == '1'

def load_data():
    w=openpyxl.load_workbook(BASE/'附件2.xlsx',data_only=True,read_only=True)
    lr=list(w['小区负载'].values); gr=list(w['光伏发电实际功率'].values)
    dates=np.array([r[0] for r in lr[1:]],dtype='datetime64[D]')
    order=np.r_[143,np.arange(143)]
    L=np.array([r[1:] for r in lr[1:]],float)[:,order]
    G=np.array([r[1:] for r in gr[1:]],float)[:,order]
    a=openpyxl.load_workbook(BASE/'附件1.xlsx',data_only=True,read_only=True).active
    price=np.array([a.cell(r,2).value for r in range(2,146)],float)[order]
    ap=list(openpyxl.load_workbook(BASE/'附件4.xlsx',data_only=True,read_only=True).active.values)
    price_day=np.array([r[1:] for r in ap[1:]],float)[:,order]
    f=list(openpyxl.load_workbook(BASE/'附件3.xlsx',data_only=True,read_only=True).active.values)
    fc={}
    for i in range(1,len(f),4):
        if f[i][0] is None: continue
        day=np.datetime64(pd.to_datetime(str(f[i][0])).date())
        for j,rel in enumerate([0,6,12,18]):
            vals=np.array(f[i+j][2:26],float)
            fc[(day,rel)]=vals
    return dates,L,G,price,price_day,fc,order

def pv_10min(day,release,fc):
    """将“预报1小时...24小时”映射到 10 min；释放时刻前不使用新预报。"""
    vals=fc[(day,release)]; out=np.zeros(T)
    for h,v in enumerate(vals,1):
        t0=release*6+h*6-6; t1=min(T,release*6+h*6)
        if t0<T: out[t0:t1]=v
    return out

def lp_plan(load,pv,price,s0,pbase=None,start=0,terminal_soc=None,net_target=None):
    # blocks P,C,D,R,S, zminus,zplus,E
    n=8*T; c=np.zeros(n)
    if pbase is None:
        c[:T]=price
    else:
        c[5*T:6*T]=.5*price; c[6*T:7*T]=1.5*price; c[7*T:8*T]=5*price
        # only future intervals are adjustable; past variables are fixed by bounds
    A=[]; b=[]
    for t in range(start,T):
        r=np.zeros(n); r[t]=1; r[T+t]=-1; r[2*T+t]=1; r[3*T+t]=-1; r[7*T+t]=1
        rhs_net=(load[t]-pv[t]) if net_target is None else net_target[t]
        A.append(r); b.append(rhs_net*DT)
        r=np.zeros(n); r[4*T+t]=1; r[T+t]=-ETA; r[2*T+t]=1/ETA
        rhs=s0 if t==start else 0
        if t>start: r[4*T+t-1]=-1
        A.append(r); b.append(rhs)
    r=np.zeros(n); r[4*T+T-1]=1; A.append(r); b.append(s0 if terminal_soc is None else terminal_soc)
    bounds=[(0,None)]*n
    for t in range(T):
        bounds[T+t]=(0,PMAX); bounds[2*T+t]=(0,PMAX); bounds[3*T+t]=(0,max(pv[t]*DT,0)); bounds[4*T+t]=(SMIN,SMAX)
        if pbase is not None:
            if t<start: bounds[t]=(pbase[t],pbase[t])
            else: bounds[t]=(0,None)
            bounds[5*T+t]=(0,None); bounds[6*T+t]=(0,None)
        if pbase is None: bounds[7*T+t]=(0,0)
    # link adjustment vars: P-P0 = zplus-zminus
    if pbase is not None:
        for t in range(start,T):
            r=np.zeros(n); r[t]=1; r[5*T+t]=1; r[6*T+t]=-1; A.append(r); b.append(pbase[t])
    res=linprog(c,A_eq=np.array(A),b_eq=np.array(b),bounds=bounds,method='highs')
    if not res.success: raise RuntimeError(res.message)
    return res.x[:T],res.x[4*T:5*T]

def copy_style(src,dst):
    dst._style=copy(src._style); dst.font=copy(src.font); dst.fill=copy(src.fill); dst.border=copy(src.border); dst.alignment=copy(src.alignment); dst.number_format=src.number_format

def event_groups(e):
    nz=np.flatnonzero(e>1e-6); out=[]
    if len(nz)==0:return out
    a=pr=nz[0]
    for i in nz[1:]:
        if i==pr+1:pr=i
        else:out.append((a,pr,float(e[a:pr+1].sum())));a=pr=i
    out.append((a,pr,float(e[a:pr+1].sum())));return out

def hhmm(i):
    m=(i*10)%1440; return f'{m//60}:{m%60:02d}'

dates,L,G,price,price_day,fc,order=load_data(); start=int(np.flatnonzero(dates==np.datetime64('2025-02-01'))[0]); records=[]; s0=6000.
for d in range(start,len(dates)):
    day=dates[d]
    pv0=pv_10min(day,0,fc)
    price_d=price_day[d]
    hist_net=np.quantile(L[max(0,d-WINDOW):d]-G[max(0,d-WINDOW):d],ALPHA,axis=0)
    net0=np.maximum(L[d]-pv0,hist_net) if RISK_AWARE else None
    day_s0=s0
    P0,plan_s=lp_plan(L[d],pv0,price_d,day_s0,terminal_soc=day_s0,net_target=net0)
    # 分段执行实际运行结果，再对尚未发生时段调整；不重复执行已发生区间。
    P=P0.copy(); cur_s=day_s0; last=0
    C=np.zeros(T); D=np.zeros(T); E=np.zeros(T); R=np.zeros(T); S=np.zeros(T)
    for rel in [6,12,18]:
        end=rel*6
        seg=execute_fixed_plan(P,L[d],G[d],price_d,cur_s,start=last,end=end,target_soc=plan_s[end-1])
        for dst,src in zip((C,D,E,R,S),seg): dst[last:end]=src
        cur_s=float(S[end-1])
        pv_rel=pv_10min(day,rel,fc)
        net_rel=np.maximum(L[d]-pv_rel,hist_net) if RISK_AWARE else None
        pnew,plan_s=lp_plan(L[d],pv_rel,price_d,cur_s,pbase=P,start=end,terminal_soc=day_s0,net_target=net_rel)
        P[end:]=pnew[end:]
        last=end
    seg=execute_fixed_plan(P,L[d],G[d],price_d,cur_s,start=last,end=T,target_soc=day_s0)
    for dst,src in zip((C,D,E,R,S),seg): dst[last:T]=src
    for arr in (P0,P,C,D,E,R,S): arr[np.abs(arr)<1e-8]=0.0
    adj=np.maximum(P0-P,0); inc=np.maximum(P-P0,0)
    plan_cost=float(price_d@P0); adj_cost=float((.5*price_d)@adj+(1.5*price_d)@inc); em_cost=float((5*price_d)@E)
    records.append({'date':day,'P0':P0,'P':P,'C':C,'D':D,'E':E,'R':R,'S':S,'s0':day_s0,'plan_cost':plan_cost,'adj_cost':adj_cost,'em_cost':em_cost})
    s0=float(S[-1])
    if (d-start+1)%30==0: print('days solved',d-start+1,str(day))

template=prepare_result('result4-3.xlsx'); wb=openpyxl.load_workbook(template); wp=wb['计划购电量']; wa=wb['调整购电量']; wc=wb['充放电量']; we=wb['紧急购电量']
for i,rec in enumerate(records,2):
    p0raw=np.empty(T); praw=np.empty(T); p0raw[order]=rec['P0']; praw[order]=rec['P']
    for j in range(T): wp.cell(i,j+2).value=float(p0raw[j]); wa.cell(i,j+2).value=float(praw[j])
    wp.cell(i,146).value=float(p0raw.sum()); wp.cell(i,147).value=float(rec['plan_cost'])
    wa.cell(i,146).value=float(praw.sum()); wa.cell(i,147).value=float(rec['plan_cost']+rec['adj_cost'])
for row in wc.iter_rows(min_row=2):
    for cell in row: cell.value=None
for i,rec in enumerate(records):
    base=2+6*i
    for g,label in enumerate(['0:00-4:00','4:00-8:00','8:00-12:00','12:00-16:00','16:00-20:00','20:00-24:00']):
        r=base+g; wc.cell(r,1).value=pd.Timestamp(str(rec['date'])).to_pydatetime() if g==0 else None; wc.cell(r,2).value=label; wc.cell(r,3).value=float(rec['C'][24*g:24*(g+1)].sum()); wc.cell(r,4).value=float(rec['D'][24*g:24*(g+1)].sum())
    wc.cell(base,5).value='0:00'; wc.cell(base,6).value=float(rec['s0']); wc.cell(base+1,5).value='24:00'; wc.cell(base+1,6).value=float(rec['S'][-1])
for row in we.iter_rows(min_row=2):
    for cell in row: cell.value=None
r=2
for rec in records:
    ev=event_groups(rec['E']) or [(None,None,None)]
    for k,(a,b,v) in enumerate(ev):
        we.cell(r,1).value=pd.Timestamp(str(rec['date'])).to_pydatetime() if k==0 else None
        if a is not None: we.cell(r,2).value=f'{hhmm(a)}-{hhmm(b+1)}'; we.cell(r,3).value=v
        r+=1
wb.save(template)
np.savez_compressed(OUT/'q4_3_results.npz',dates=np.array([r['date'] for r in records]),P0=np.array([r['P0'] for r in records]),P=np.array([r['P'] for r in records]),C=np.array([r['C'] for r in records]),D=np.array([r['D'] for r in records]),E=np.array([r['E'] for r in records]),S=np.array([r['S'] for r in records]),plan_cost=np.array([r['plan_cost'] for r in records]),adj_cost=np.array([r['adj_cost'] for r in records]),em_cost=np.array([r['em_cost'] for r in records]))
print('saved',template,'risk_aware',RISK_AWARE,'days',len(records),'plan',sum(r['plan_cost'] for r in records),'adjust',sum(r['adj_cost'] for r in records),'emergency_kwh',sum(r['E'].sum() for r in records),'emergency_cost',sum(r['em_cost'] for r in records),'total',sum(r['plan_cost']+r['adj_cost']+r['em_cost'] for r in records))
for target in ['2025-03-20','2025-06-21','2025-09-23','2025-12-21']:
    rec=next(x for x in records if str(x['date'])==target); print(target,'P0',rec['P0'].sum(),'Pfinal',rec['P'].sum(),'adjcost',rec['adj_cost'],'emergency',event_groups(rec['E']))
