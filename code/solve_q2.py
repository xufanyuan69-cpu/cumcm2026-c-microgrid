from pathlib import Path
from copy import copy
import numpy as np
import openpyxl
from scipy.optimize import milp, Bounds, LinearConstraint
from scipy.sparse import lil_matrix, vstack
from dispatch_execution import execute_fixed_plan

from project_paths import input_directory, analysis_directory, prepare_result

BASE=input_directory()
OUT=analysis_directory()
DT=1/6; ETA=.9; PMAX=5000*DT; SMIN=1200.; SMAX=10800.; T=144
ALPHA=.90; WINDOW=30

def load_inputs():
    w=openpyxl.load_workbook(BASE/'附件2.xlsx',data_only=True,read_only=True)
    lr=list(w['小区负载'].values); gr=list(w['光伏发电实际功率'].values)
    dates=np.array([r[0] for r in lr[1:]],dtype='datetime64[D]')
    Lraw=np.array([r[1:] for r in lr[1:]],float); Graw=np.array([r[1:] for r in gr[1:]],float)
    # 与 Q1 相同：把每行末尾“次日0:00”作为重复日0:00移到首位。
    order=np.r_[T-1,np.arange(T-1)]
    L=Lraw[:,order]; G=Graw[:,order]
    a=openpyxl.load_workbook(BASE/'附件1.xlsx',data_only=True,read_only=True).active
    praw=np.array([a.cell(r,2).value for r in range(2,146)],float)
    price=praw[order]
    return dates,L,G,price,order

def idx(b,t): return b*T+t

def solve_plan(net,price,s0):
    # blocks: P,C,D,R,S,yc,yd,w
    n=8*T; c=np.zeros(n)
    c[:T]=price; c[T:3*T]=1e-7; c[7*T:8*T]=1e-4
    Aeq=lil_matrix((2*T+1,n)); beq=np.zeros(2*T+1)
    for t in range(T):
        Aeq[2*t,idx(0,t)]=1; Aeq[2*t,idx(1,t)]=-1; Aeq[2*t,idx(2,t)]=1; Aeq[2*t,idx(3,t)]=-1
        beq[2*t]=net[t]
        Aeq[2*t+1,idx(4,t)]=1; Aeq[2*t+1,idx(1,t)]=-ETA; Aeq[2*t+1,idx(2,t)]=1/ETA
        if t: Aeq[2*t+1,idx(4,t-1)]=-1
        else: beq[2*t+1]=s0
    Aeq[-1,idx(4,T-1)]=1; beq[-1]=s0
    Aub=lil_matrix((5*T,n)); bub=np.zeros(5*T)
    for t in range(T):
        Aub[5*t,idx(1,t)]=1; Aub[5*t,idx(5,t)]=-PMAX
        Aub[5*t+1,idx(2,t)]=1; Aub[5*t+1,idx(6,t)]=-PMAX
        Aub[5*t+2,idx(5,t)]=1; Aub[5*t+2,idx(6,t)]=1; bub[5*t+2]=1
        # |(yc-yd)_t-(yc-yd)_{t-1}| <= w_t
        Aub[5*t+3,idx(5,t)]=1; Aub[5*t+3,idx(6,t)]=-1; Aub[5*t+3,idx(7,t)]=-1
        Aub[5*t+4,idx(5,t)]=-1; Aub[5*t+4,idx(6,t)]=1; Aub[5*t+4,idx(7,t)]=-1
        if t:
            Aub[5*t+3,idx(5,t-1)]=-1; Aub[5*t+3,idx(6,t-1)]=1
            Aub[5*t+4,idx(5,t-1)]=1; Aub[5*t+4,idx(6,t-1)]=-1
    lb=np.zeros(n); ub=np.full(n,np.inf)
    ub[T:3*T]=PMAX; ub[3*T:4*T]=np.maximum(-net,0)
    lb[4*T:5*T]=SMIN; ub[4*T:5*T]=SMAX
    ub[5*T:7*T]=1; ub[7*T:8*T]=2
    integrality=np.zeros(n,int); integrality[5*T:7*T]=1
    cons=[LinearConstraint(Aeq.tocsr(),beq,beq),LinearConstraint(Aub.tocsr(),-np.inf,bub)]
    res=milp(c,integrality=integrality,bounds=Bounds(lb,ub),constraints=cons,
             options={'time_limit':20,'mip_rel_gap':1e-7})
    if not res.success: raise RuntimeError(res.message)
    x=res.x
    return x[:T],x[T:2*T],x[2*T:3*T],x[4*T:5*T]

def copy_style(src,dst):
    if src.has_style:
        dst._style=copy(src._style)
    dst.number_format=src.number_format; dst.alignment=copy(src.alignment)
    dst.font=copy(src.font); dst.fill=copy(src.fill); dst.border=copy(src.border)

def group_events(e):
    nz=np.flatnonzero(e>1e-6); out=[]
    if len(nz)==0: return out
    a=prev=nz[0]
    for i in nz[1:]:
        if i==prev+1: prev=i
        else: out.append((a,prev,float(e[a:prev+1].sum()))); a=prev=i
    out.append((a,prev,float(e[a:prev+1].sum())))
    return out

def hhmm(step):
    mins=(step*10)%(24*60); return f'{mins//60}:{mins%60:02d}'

dates,L,G,price,order=load_inputs(); start=int(np.flatnonzero(dates==np.datetime64('2025-02-01'))[0])
s0=6000.; records=[]
for d in range(start,len(dates)):
    hist=(L[max(0,d-WINDOW):d]-G[max(0,d-WINDOW):d])*DT
    net=np.quantile(hist,ALPHA,axis=0)
    P,pc,pd,ps=solve_plan(net,price,s0)
    C,D,E,R,S=execute_fixed_plan(P,L[d],G[d],price,s0,target_soc=s0)
    for arr in (P,C,D,E,R,S): arr[np.abs(arr)<1e-8]=0.0
    records.append({'date':dates[d],'P':P,'C':C,'D':D,'E':E,'R':R,'S':S,'s0':s0,
                    'plan_cost':float(price@P),'emergency_cost':float((5*price)@E)})
    s0=float(S[-1])
    if (d-start+1)%30==0: print('days solved',d-start+1,'last',str(dates[d]))

template=prepare_result('result2.xlsx'); wb=openpyxl.load_workbook(template)
wp=wb['计划购电量']; wc=wb['充放电量']; we=wb['紧急购电量']
for i,rec in enumerate(records,2):
    p_raw=np.empty(T); p_raw[order]=rec['P']
    for j,v in enumerate(p_raw,2): wp.cell(i,j).value=float(v)
    wp.cell(i,146).value=float(p_raw.sum()); wp.cell(i,147).value=float(rec['plan_cost'])

# 重建完整充放电明细表。
for row in wc.iter_rows(min_row=2):
    for cell in row: cell.value=None
for i,rec in enumerate(records):
    base=2+6*i
    for r in range(base,base+6):
        for col in range(1,7): copy_style(wc.cell(2+(r-base),col),wc.cell(r,col))
    wc.cell(base,1).value=np.datetime64(rec['date']).astype('datetime64[D]').astype(object)
    for g,label in enumerate(['0:00-4:00','4:00-8:00','8:00-12:00','12:00-16:00','16:00-20:00','20:00-24:00']):
        r=base+g; wc.cell(r,2).value=label; wc.cell(r,3).value=float(rec['C'][24*g:24*(g+1)].sum()); wc.cell(r,4).value=float(rec['D'][24*g:24*(g+1)].sum())
    wc.cell(base,5).value='0:00'; wc.cell(base,6).value=float(rec['s0'])
    wc.cell(base+1,5).value='24:00'; wc.cell(base+1,6).value=float(rec['S'][-1])

# 重建紧急购电表：每日至少保留一行；有多段时逐段列出。
for row in we.iter_rows(min_row=2):
    for cell in row: cell.value=None
r=2
for rec in records:
    events=group_events(rec['E']) or [(None,None,None)]
    for k,(a,b,val) in enumerate(events):
        for col in range(1,4): copy_style(we.cell(2,col),we.cell(r,col))
        if k==0: we.cell(r,1).value=np.datetime64(rec['date']).astype('datetime64[D]').astype(object)
        if a is not None:
            we.cell(r,2).value=f'{hhmm(a)}-{hhmm(b+1)}'; we.cell(r,3).value=val
        r+=1

wb.save(template)
np.savez_compressed(OUT/'q2_results.npz',dates=np.array([r['date'] for r in records]),P=np.array([r['P'] for r in records]),C=np.array([r['C'] for r in records]),D=np.array([r['D'] for r in records]),E=np.array([r['E'] for r in records]),R=np.array([r['R'] for r in records]),S=np.array([r['S'] for r in records]),s0=np.array([r['s0'] for r in records]),plan_cost=np.array([r['plan_cost'] for r in records]),emergency_cost=np.array([r['emergency_cost'] for r in records]))
print('saved',template,'days',len(records),'plan cost',sum(r['plan_cost'] for r in records),'emergency kWh',sum(r['E'].sum() for r in records),'emergency cost',sum(r['emergency_cost'] for r in records))
for target in ['2025-03-20','2025-06-21','2025-09-23','2025-12-21']:
    rec=next(x for x in records if str(x['date'])==target); print(target,group_events(rec['E']))
